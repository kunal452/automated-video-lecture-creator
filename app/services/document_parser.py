import re
from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
from pptx import Presentation

def parse_pdf(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extracts text from each PDF page and heuristically identifies
    slide titles and bullet points.
    """
    reader = PdfReader(str(file_path))
    slides = []

    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        if not lines:
            continue

        # Heuristic: First prominent line is often title
        title = lines[0]
        bullets = []
        raw_body = []

        for line in lines[1:]:
            # Check for bullet symbols or numbered points
            clean_line = re.sub(r"^[\*\-\•\d+\.]\s*", "", line).strip()
            if len(clean_line) > 3:
                bullets.append(clean_line)
                raw_body.append(clean_line)

        # Default script: conversational expansion
        script = f"In slide {idx}, we explore {title}. " + " ".join(raw_body[:4])
        if len(bullets) == 0:
            bullets = ["Key concept overview", "Detailed discussion"]

        slides.append({
            "index": idx,
            "title": title[:80],
            "bullet_points": bullets[:5],
            "code_snippet": None,
            "spoken_script": script
        })

    return slides

def parse_pptx(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extracts structured slides from a PowerPoint presentation,
    retaining slide titles, bullet points, and speaker notes.
    """
    prs = Presentation(str(file_path))
    slides = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = f"Slide {idx}"
        bullets = []
        speaker_notes = ""

        # Extract title shape if present
        if slide.shapes.title and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()

        # Extract bullet text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = paragraph.text.strip()
                    if p_text and p_text != title and len(p_text) > 3:
                        bullets.append(p_text)

        # Extract presenter notes if available
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            speaker_notes = slide.notes_slide.notes_text_frame.text.strip()

        # Generate spoken script from notes or bullets
        if speaker_notes:
            spoken_script = speaker_notes
        else:
            bullets_narration = ". ".join(bullets[:4]) if bullets else "Let us explore this topic in detail."
            spoken_script = f"Welcome to slide {idx}: {title}. {bullets_narration}"

        if not bullets:
            bullets = ["Core principle overview", "Practical applications and implications"]

        slides.append({
            "index": idx,
            "title": title[:80],
            "bullet_points": bullets[:5],
            "code_snippet": None,
            "spoken_script": spoken_script
        })

    return slides

def parse_text_notes(text: str) -> List[Dict[str, Any]]:
    """
    Parses Markdown or plain text notes delimited by '# ' or '---'.
    """
    chunks = re.split(r"(?m)^(?=#{1,2}\s+|\-\-\-)", text)
    slides = []
    idx = 1

    for chunk in chunks:
        chunk = chunk.strip().lstrip("-").strip()
        if not chunk:
            continue

        lines = [l.strip() for l in chunk.split("\n") if l.strip()]
        if not lines:
            continue

        first_line = lines[0].lstrip("#").strip()
        title = first_line if first_line else f"Topic {idx}"
        bullets = []
        body_lines = []

        for line in lines[1:]:
            clean = re.sub(r"^[\*\-\•\d+\.]\s*", "", line).strip()
            if clean:
                bullets.append(clean)
                body_lines.append(clean)

        script = f"Here we discuss {title}. " + ". ".join(body_lines[:3])
        if not bullets:
            bullets = ["Primary takeaway", "Detailed considerations"]

        slides.append({
            "index": idx,
            "title": title[:80],
            "bullet_points": bullets[:5],
            "code_snippet": None,
            "spoken_script": script
        })
        idx += 1

    return slides
